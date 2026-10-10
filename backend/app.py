from flask import Flask, request, send_file
from flask_cors import CORS
import io
from docx import Document

app = Flask(__name__)
CORS(app)

@app.route('/generate', methods=['POST'])
def generate_document():
    try:
        data = request.get_json()
        
        # Extract form fields sent from React frontend
        full_name = data.get('fullName', '')
        recipient = data.get('recipient', '')
        date = data.get('date', '')
        amount = data.get('amount', '')
        subject = data.get('subject', '')
        issue_description = data.get('issueDescription', '')

        # Create a new Word document using python-docx
        doc = Document()
        doc.add_heading('Draft Letter', 0)
        
        doc.add_paragraph(f"Date: {date}")
        doc.add_paragraph(f"To:\n{recipient}")
        doc.add_paragraph(f"Subject: {subject}")
        
        doc.add_paragraph("Respected Sir/Madam,")
        doc.add_paragraph(f"I am writing to request your attention to the following matter:\n{issue_description}")
        doc.add_paragraph(f"Amount involved: {amount}")
        doc.add_paragraph("I request that you review this matter and take appropriate action.")
        doc.add_paragraph("Thank you.")
        doc.add_paragraph(f"Sincerely,\n{full_name}")

        # Save document to a BytesIO stream
        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)

        return send_file(
            file_stream,
            as_attachment=True,
            download_name='rightspocket-draft-letter.docx',
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
    except Exception as e:
        print("Error generating document:", str(e))
        return {"error": str(e)}, 500

if __name__ == '__main__':
    app.run(debug=True, port=5000)